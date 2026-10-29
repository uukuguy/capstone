"""Verify the isolated programming environment and emit a safe build receipt.

No Provider, network request, model, Authority, or user data is used.
"""
import argparse
import hashlib
import importlib
from importlib import metadata
import json
from pathlib import Path
import platform
import re
import subprocess
import tempfile
import warnings


SCHEMA = 'capstone-native-environment/1'


def command(args, *, cwd=None):
    result = subprocess.run(args, cwd=cwd, capture_output=True, text=True, timeout=30)
    if result.returncode:
        raise RuntimeError('native environment command check failed: ' + args[0])
    return result.stdout + result.stderr


def inspect_inventory(manifest, lock_text):
    packages = {}
    for item in manifest['packages']:
        name = item['requirement'].split('==')[0]
        importlib.import_module(item['module'])
        version = metadata.version(name)
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9.+!_-]{0,80}', version):
            raise ValueError('invalid installed package version')
        normalized = re.sub(r'[-_.]+', '-', name).lower()
        locked = {match[1] for match in re.findall(r'^([A-Za-z0-9_.-]+)==([^\s;]+)', lock_text, re.M)
                  if re.sub(r'[-_.]+', '-', match[0]).lower() == normalized}
        if version not in locked:
            raise ValueError('installed package does not match native analysis lock: ' + name)
        packages[name] = version
    commands = {}
    for item in manifest['commands']:
        output = command([item['name'], *item['version_args']])
        match = re.search(item['version_pattern'], output)
        if match is None or not re.fullmatch(r'[0-9.]{1,40}', match[1]):
            raise ValueError('command version could not be verified: ' + item['name'])
        commands[item['name']] = match[1]
    return packages, commands


def exercise_report(directory, *, renderer=None):
    """Check actual CLI, exact CSV strings and inline figures, without a model."""
    (directory / 'report-values.csv').write_text('label,value,ratio\n中文,001.2300,12.50%\n', encoding='utf-8')
    (directory / 'report-figure.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><rect width="20" height="20" fill="teal"/></svg>')
    spec = {'title': '中文报告校验', 'sections': [{'id': 'results', 'title': '结果与限制', 'blocks': [
        {'type': 'paragraph', 'text': '保存的数值 < 不作重算'},
        {'type': 'table', 'csv': 'report-values.csv'},
        {'type': 'figure', 'path': 'report-figure.svg', 'caption': '实际保存的示意图'}]}]}
    if (directory / 'chart.svg').is_file():
        spec['sections'][0]['blocks'].append({'type': 'figure', 'path': 'chart.svg', 'caption': '实际中文统计图'})
    (directory / 'report-spec.json').write_text(json.dumps(spec, ensure_ascii=False), encoding='utf-8')
    receipt = json.loads(command([*(renderer or ['capstone-report']), '--spec', 'report-spec.json', '--output', 'report.html'], cwd=directory))
    assert receipt['sections'] == 1
    assert receipt['figures'] == sum(block['type'] == 'figure' for block in spec['sections'][0]['blocks'])
    for source, digest in receipt['source_sha256'].items():
        assert hashlib.sha256((directory / source).read_bytes()).hexdigest() == digest
    rendered = (directory / 'report.html').read_text()
    assert all(value in rendered for value in ('001.2300', '12.50%', 'data:image/svg+xml;base64,', '图 1 · 实际保存的示意图', 'href="#section-results"', '&lt;'))
    assert '<script' not in rendered and 'src="http' not in rendered


def exercise_presentation(directory, *, renderer=None):
    """Verify deterministic PPTX structure, literal data, figures and input hashes."""
    from PIL import Image
    from pptx import Presentation
    from zipfile import ZipFile
    (directory/'slide-values.csv').write_text('label,value,ratio\n中文,001.2300,12.50%\n',encoding='utf-8')
    Image.new('RGB',(320,180),'teal').save(directory/'slide-figure.png')
    spec={'schema':'capstone-presentation/1','title':'中文演示校验','slides':[{
        'title':'结果与限制','source_notes':['保存的输入；不验证结论'], 'blocks':[
            {'type':'paragraph','text':'保存的数值 < 不作重算'},
            {'type':'table','csv':'slide-values.csv','max_rows':10},
            {'type':'figure','path':'slide-figure.png','caption':'实际保存的示意图'}]}]}
    (directory/'presentation-spec.json').write_text(json.dumps(spec,ensure_ascii=False),encoding='utf-8')
    receipt=json.loads(command([*(renderer or ['capstone-report']),'--spec','presentation-spec.json',
        '--output','presentation.pptx'],cwd=directory))
    assert receipt['schema']=='capstone-presentation-receipt/1'
    deck=Presentation(directory/'presentation.pptx')
    assert len(deck.slides)==receipt['slides'] and len(deck.slides)>=4
    text='\n'.join(shape.text if shape.has_text_frame else '\n'.join(cell.text for row in shape.table.rows for cell in row.cells)
        if shape.has_table else '' for slide in deck.slides for shape in slide.shapes)
    assert all(value in text for value in ['001.2300','12.50%','保存的数值','实际保存的示意图'])
    assert sum(shape.shape_type==13 for slide in deck.slides for shape in slide.shapes)==1
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=deck.slide_width+100 and shape.top+shape.height<=deck.slide_height+100
    for source,digest in receipt['source_sha256'].items():
        assert hashlib.sha256((directory/source).read_bytes()).hexdigest()==digest
    with ZipFile(directory/'presentation.pptx') as archive:
        assert not any(b'TargetMode="External"' in archive.read(path) for path in archive.namelist() if path.endswith('.rels'))


def exercise(directory):
    import numpy as np
    import pandas as pd
    from scipy import stats
    import statsmodels.api as sm
    from sklearn.linear_model import LinearRegression
    import matplotlib
    matplotlib.use('Agg', force=True)
    from matplotlib import pyplot as plt, font_manager
    import seaborn as sns
    import plotly.express as px
    from PIL import Image
    import openpyxl
    import duckdb
    from pypdf import PdfWriter, PdfReader
    import pdfplumber
    from docx import Document
    from pptx import Presentation
    from bs4 import BeautifulSoup
    from lxml import etree
    from jinja2 import Environment
    import requests
    import httpx
    import yaml
    from tabulate import tabulate

    verified = []
    frame = pd.DataFrame({'x': [1., 2., 3.], 'y': [2., 4., 6.]})
    frame.to_csv(directory / 'data.csv', index=False)
    pd.testing.assert_frame_equal(frame, pd.read_csv(directory / 'data.csv'))
    verified.append('csv_dataframes')
    assert np.isclose(frame.x.corr(pd.Series([2., 3., 1.]), method='spearman'), -0.5)
    assert np.isclose(stats.spearmanr(frame.x, frame.y).statistic, 1.)
    assert np.isclose(sm.OLS(frame.y, sm.add_constant(frame.x)).fit().params['x'], 2.)
    verified.append('statistics')
    assert np.allclose(LinearRegression().fit(frame[['x']], frame.y).predict(pd.DataFrame({'x': [4.]})), [8.])
    verified.append('machine_learning')
    font_manager.findfont('Noto Sans CJK SC', fallback_to_default=False)
    with warnings.catch_warnings():
        warnings.filterwarnings('error', message='Glyph .* missing from font')
        with plt.rc_context({'font.family': 'Noto Sans CJK SC'}):
            fig, ax = plt.subplots()
            sns.lineplot(data=frame, x='x', y='y', ax=ax)
            ax.set_title('中文图表校验')
            fig.savefig(directory / 'chart.svg')
            plt.close(fig)
    etree.parse(str(directory / 'chart.svg'))
    verified.append('chinese_plots')
    px.line(frame, x='x', y='y', title='中文交互图').write_html(directory / 'interactive.html', include_plotlyjs=True)
    html = (directory / 'interactive.html').read_text()
    assert 'Plotly.newPlot' in html
    assert not BeautifulSoup(html, 'html.parser').find_all('script', src=True)
    verified.append('interactive_html')
    Image.new('RGB', (8, 8), 'blue').save(directory / 'image.png')
    with Image.open(directory / 'image.png') as image:
        assert image.size == (8, 8)
    verified.append('images')
    for engine in ('openpyxl', 'xlsxwriter'):
        path = directory / (engine + '.xlsx')
        frame.to_excel(path, index=False, engine=engine)
        workbook = openpyxl.load_workbook(path, read_only=True)
        assert workbook.active['B4'].value == 6
        workbook.close()
    verified.append('excel')
    frame.to_parquet(directory / 'data.parquet', index=False, engine='pyarrow')
    pd.testing.assert_frame_equal(frame, pd.read_parquet(directory / 'data.parquet'))
    verified.append('parquet')
    connection = duckdb.connect(':memory:')
    try:
        connection.register('sample', frame)
        assert connection.execute('select sum(y) from sample').fetchone()[0] == 12.
    finally:
        connection.close()
    verified.append('sql')
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(directory / 'sample.pdf')
    assert len(PdfReader(directory / 'sample.pdf').pages) == 1
    with pdfplumber.open(directory / 'sample.pdf') as pdf:
        assert pdf.pages[0].width == 72
    verified.append('pdf')
    document = Document()
    document.add_paragraph('中文文档校验')
    document.save(directory / 'sample.docx')
    assert Document(directory / 'sample.docx').paragraphs[0].text == '中文文档校验'
    verified.append('word')
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000000, 500000).text = '中文演示校验'
    presentation.save(directory / 'sample.pptx')
    assert Presentation(directory / 'sample.pptx').slides[0].shapes[0].text == '中文演示校验'
    verified.append('powerpoint')
    rendered = Environment(autoescape=True).from_string('<p>{{ text }}</p>').render(text='中文<&')
    assert BeautifulSoup(rendered, 'html.parser').p.text == '中文<&'
    etree.fromstring(rendered.encode())
    verified.append('html')
    # Construct requests only. The smoke must never contact an external server.
    assert requests.Request('GET', 'https://example.invalid/').prepare().method == 'GET'
    assert httpx.Request('GET', 'https://example.invalid/').method == 'GET'
    verified.append('http_clients')
    assert yaml.safe_load(yaml.safe_dump({'items': [1, 2]})) == {'items': [1, 2]}
    verified.append('yaml')
    assert 'value' in tabulate([['value', 3]])
    verified.append('tables')
    (directory / 'search.txt').write_text('capstone smoke\n')
    assert 'capstone' in command(['rg', 'capstone', 'search.txt'], cwd=directory)
    (directory / 'data.json').write_text('{"count":3}')
    assert command(['jq', '.count', 'data.json'], cwd=directory).strip() == '3'
    command(['zip', '-q', 'sample.zip', 'search.txt'], cwd=directory)
    assert command(['unzip', '-p', 'sample.zip', 'search.txt'], cwd=directory) == 'capstone smoke\n'
    verified.append('archive_cli')
    command(['pdftotext', 'sample.pdf', 'sample.txt'], cwd=directory)
    command(['pdftoppm', '-singlefile', '-scale-to', '32', '-png', 'sample.pdf', 'rendered'], cwd=directory)
    with Image.open(directory / 'rendered.png') as image:
        assert image.size == (32, 32)
    verified.append('pdf_cli')
    assert command(['node', '-e', 'console.log(2 + 3)']).strip() == '5'
    verified.append('javascript')
    (directory / 'sample.ts').write_text('const value: number = 5; console.log(value);\n')
    command(['tsc', '--target', 'ES2022', '--skipLibCheck', '--outDir', 'compiled', 'sample.ts'], cwd=directory)
    assert command(['node', 'compiled/sample.js'], cwd=directory).strip() == '5'
    assert command(['tsx', 'sample.ts'], cwd=directory).strip() == '5'
    verified.append('typescript')
    exercise_report(directory)
    verified.append('standalone_reports')
    exercise_presentation(directory)
    verified.append('standalone_presentations')
    return verified


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--manifest', type=Path, default=Path('/opt/general/configs/runtime/native-analysis.json'))
    parser.add_argument('--lock', type=Path, default=Path('/opt/general/configs/runtime/native-analysis.lock.txt'))
    parser.add_argument('--inventory-output', type=Path, default=Path('/opt/general/runtime-environment.json'))
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_text())
    if manifest['schema'] != 'capstone-native-analysis/1':
        raise ValueError('unsupported native analysis manifest')
    packages, commands = inspect_inventory(manifest, args.lock.read_text())
    with tempfile.TemporaryDirectory(prefix='native-smoke-') as temporary:
        verified = exercise(Path(temporary))
    receipt = {'schema': SCHEMA,
        'python': {'version': platform.python_version(), 'implementation': platform.python_implementation()},
        'packages': packages, 'commands': commands, 'smoke_capabilities': verified,
        'manifest_sha256': hashlib.sha256(args.manifest.read_bytes()).hexdigest(),
        'lock_sha256': hashlib.sha256(args.lock.read_bytes()).hexdigest()}
    # Only a complete successful smoke creates or replaces the trusted receipt.
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=args.inventory_output.parent, delete=False) as output:
        temporary = Path(output.name)
        json.dump(receipt, output, ensure_ascii=False, sort_keys=True)
        output.write('\n')
    temporary.replace(args.inventory_output)
    print(json.dumps({'schema': SCHEMA, 'verified_capabilities': len(verified), 'packages': len(packages), 'commands': len(commands)}))


if __name__ == '__main__':
    main()
