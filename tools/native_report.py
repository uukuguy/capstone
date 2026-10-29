#!/usr/bin/env python3
"""Render saved report text, CSV strings and figures without interpreting facts."""
import argparse
import base64
import csv
import hashlib
from html import escape
import io
import json
import os
from pathlib import Path
import re
import secrets
import stat
import sys
import unicodedata
from xml.etree import ElementTree


VERSION = '1.1.0'
CSS = '''
:root{color-scheme:light;font-family:system-ui,-apple-system,"Noto Sans CJK SC",sans-serif;color:#172b3b;background:#edf2f5}
*{box-sizing:border-box}body{margin:0;line-height:1.75}main{max-width:1100px;margin:32px auto;background:white;padding:48px 56px;border:1px solid #d8e2e8;border-radius:12px}
header{border-bottom:3px solid #217778;padding-bottom:24px;margin-bottom:28px}h1{font-size:2rem;line-height:1.35;margin:0 0 12px}h2{font-size:1.35rem;color:#175e61;margin:36px 0 16px}header p{color:#526774}
nav{background:#f3f7f8;border:1px solid #d8e2e8;padding:16px 24px}nav a{color:#175e61;text-decoration:none}section{scroll-margin-top:24px}p{white-space:pre-wrap;overflow-wrap:anywhere}li{margin:6px 0}
.table-wrap{overflow-x:auto;margin:24px 0}table{width:100%;border-collapse:collapse;font-size:.9rem}caption{text-align:left;font-weight:600;padding:0 0 10px}th,td{border:1px solid #d8e2e8;text-align:left;vertical-align:top;padding:9px 12px;white-space:pre-wrap;overflow-wrap:anywhere}th{background:#eaf2f3}tbody tr:nth-child(even){background:#f8fafb}
figure{margin:28px 0;break-inside:avoid}figure img{display:block;max-width:100%;height:auto;margin:auto}figcaption{text-align:center;color:#526774;font-size:.92rem;margin-top:12px}footer{border-top:1px solid #d8e2e8;color:#526774;font-size:.8rem;margin-top:40px;padding-top:16px}
@media(max-width:700px){main{margin:0;padding:24px 18px;border:0;border-radius:0}h1{font-size:1.7rem}}
@media print{@page{size:A4;margin:18mm}body{background:white}main{max-width:none;margin:0;padding:0;border:0}nav{break-inside:avoid}h2,caption{break-after:avoid}thead{display:table-header-group}tr{break-inside:avoid}.table-wrap{overflow:visible}a{color:inherit;text-decoration:none}}
'''


class ReportError(ValueError):
    def __init__(self, message, *, diagnostic=None):
        super().__init__(message)
        self.diagnostic = diagnostic


def text(value, field):
    if not isinstance(value, str) or not value.strip():
        raise ReportError(field + ' must be nonempty text')
    return value


def fields(value, required, optional=()):
    if not isinstance(value, dict) or not set(required) <= value.keys() or set(value) - set(required) - set(optional):
        raise ReportError('missing or unknown report fields')


class Workspace:
    """Use directory descriptors so path checks and actual IO share a boundary."""
    def __init__(self):
        self.root = os.open('.', os.O_RDONLY | os.O_DIRECTORY)
        self.sources = {}
        self.identities = set()

    def parent(self, value):
        text(value, 'file path')
        parts = value.split('/')
        if any(part in {'', '.', '..'} for part in parts) or any(c in value for c in ('\\', ':', '\0')):
            raise ReportError('file paths must stay within the current workspace')
        current = os.dup(self.root)
        try:
            for part in parts[:-1]:
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=current)
                os.close(current)
                current = child
            return current, parts[-1]
        except BaseException:
            os.close(current)
            raise

    def read(self, value, *, location='spec', maximum=None):
        parent = None
        try:
            parent, name = self.parent(value)
            descriptor = os.open(name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=parent)
            with os.fdopen(descriptor, 'rb') as stream:
                info = os.fstat(stream.fileno())
                if not stat.S_ISREG(info.st_mode):
                    raise ReportError('report sources must be regular files')
                if info.st_nlink != 1:
                    raise ReportError('hardlinked report sources are not supported')
                if maximum is not None and info.st_size > maximum:
                    raise ReportError('report source exceeds its rendering byte budget')
                data = stream.read() if maximum is None else stream.read(maximum + 1)
                if maximum is not None and len(data) > maximum:
                    raise ReportError('report source exceeds its rendering byte budget')
                self.identities.add((info.st_dev, info.st_ino))
        except OSError as error:
            code = 'source_missing' if isinstance(error, FileNotFoundError) else 'source_unavailable'
            raise ReportError('Cannot read report source at ' + location + ': ' + json.dumps(value, ensure_ascii=False),
                diagnostic={'code': code, 'path': value, 'location': location}) from None
        finally:
            if parent is not None:
                os.close(parent)
        digest = hashlib.sha256(data).hexdigest()
        if value in self.sources and self.sources[value] != digest:
            raise ReportError('a report source changed during rendering')
        self.sources[value] = digest
        return data

    def write(self, value, data):
        parent, name = self.parent(value)
        temporary = '.capstone-report-' + secrets.token_hex(16)
        try:
            if value in self.sources:
                raise ReportError('output must not overwrite the spec or any source')
            try:
                info = os.stat(name, dir_fd=parent, follow_symlinks=False)
            except FileNotFoundError:
                info = None
            if info is not None and (not stat.S_ISREG(info.st_mode) or (info.st_dev, info.st_ino) in self.identities):
                raise ReportError('output must not overwrite a source or symbolic link')
            descriptor = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600, dir_fd=parent)
            with os.fdopen(descriptor, 'wb') as stream:
                stream.write(data)
            os.replace(temporary, name, src_dir_fd=parent, dst_dir_fd=parent)
        finally:
            try:
                os.unlink(temporary, dir_fd=parent)
            except FileNotFoundError:
                pass
            os.close(parent)


def figure_data(path, data):
    suffix = Path(path).suffix.lower()
    if suffix == '.svg':
        if b'<!ENTITY' in data.upper():
            raise ReportError('SVG entities are not supported')
        try:
            root = ElementTree.fromstring(data)
        except ElementTree.ParseError:
            raise ReportError('invalid SVG figure') from None
        if root.tag.rsplit('}', 1)[-1] != 'svg':
            raise ReportError('figure is not an SVG document')
        for node in root.iter():
            if node.tag.rsplit('}', 1)[-1].lower() in {'script', 'foreignobject', 'animate', 'animatemotion', 'animatetransform', 'set', 'discard'}:
                raise ReportError('active SVG content is not supported')
            for key, value in node.attrib.items():
                attribute = key.rsplit('}', 1)[-1].lower()
                if attribute.startswith('on'):
                    raise ReportError('active SVG attributes are not supported')
                if attribute in {'href', 'src'} and not (value.startswith('#') or re.fullmatch(r'data:image/(?:png|jpeg);base64,[A-Za-z0-9+/=\s]+', value)):
                    raise ReportError('SVG external references are not supported')
            styles = ' '.join([*node.attrib.values(), node.text or ''])
            if '@import' in styles.lower() or any(not url.strip(' \t\r\n\"\x27').startswith('#') for url in re.findall(r'url\s*\((.*?)\)', styles, re.I | re.S)):
                raise ReportError('SVG external styles are not supported')
        media = 'image/svg+xml'
    elif suffix == '.png' and data.startswith(b'\x89PNG\r\n\x1a\n'):
        media = 'image/png'
    elif suffix in {'.jpg', '.jpeg'} and data.startswith(b'\xff\xd8\xff'):
        media = 'image/jpeg'
    else:
        raise ReportError('figures must be saved PNG, JPEG or safe SVG files')
    return 'data:' + media + ';base64,' + base64.b64encode(data).decode('ascii')


def render(spec_path, output):
    if Path(output).suffix.lower() == '.pptx':
        return render_presentation(spec_path, output)
    workspace = Workspace()
    try:
        spec = json.loads(workspace.read(spec_path))
        fields(spec, ('title', 'sections'), ('subtitle', 'schema'))
        if spec.get('schema', 'capstone-report/1') != 'capstone-report/1':
            raise ReportError('unsupported report schema')
        title = escape(text(spec['title'], 'title'))
        subtitle = '<p>' + escape(text(spec['subtitle'], 'subtitle')) + '</p>' if 'subtitle' in spec else ''
        if not isinstance(spec['sections'], list) or not spec['sections']:
            raise ReportError('sections must be a nonempty list')
        identities, navigation, sections, figures = set(), [], [], 0
        for number, section in enumerate(spec['sections'], 1):
            fields(section, ('id', 'title', 'blocks'))
            identity = text(section['id'], 'section id')
            if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_-]*', identity) or identity in identities:
                raise ReportError('section ids must be unique letters, digits, hyphens or underscores')
            identities.add(identity)
            heading = escape(text(section['title'], 'section title'))
            anchor = 'section-' + identity
            navigation.append(f'<li><a href="#{anchor}">{number}. {heading}</a></li>')
            if not isinstance(section['blocks'], list):
                raise ReportError('blocks must be a list')
            blocks = []
            for block_index, block in enumerate(section['blocks']):
                location = f'sections[{number - 1}].blocks[{block_index}]'
                if not isinstance(block, dict):
                    raise ReportError('each block must be an object')
                kind = block.get('type')
                if kind == 'paragraph':
                    fields(block, ('type', 'text'))
                    blocks.append('<p>' + escape(text(block['text'], 'paragraph')) + '</p>')
                elif kind == 'bullets':
                    fields(block, ('type', 'items'))
                    if not isinstance(block['items'], list) or not block['items']:
                        raise ReportError('bullet items must be a nonempty list')
                    blocks.append('<ul>' + ''.join('<li>' + escape(text(item, 'bullet')) + '</li>' for item in block['items']) + '</ul>')
                elif kind == 'table':
                    fields(block, ('type', 'csv'), ('columns', 'caption'))
                    csv.field_size_limit(sys.maxsize)
                    reader = csv.reader(io.StringIO(workspace.read(block['csv'], location=location + '.csv').decode('utf-8-sig'), newline=''), strict=True)
                    header = next(reader, [])
                    if not header or len(set(header)) != len(header):
                        raise ReportError('CSV must have a nonempty unique header')
                    columns = block.get('columns', header)
                    if not isinstance(columns, list) or not columns or any(not isinstance(column, str) for column in columns) or len(set(columns)) != len(columns) or any(column not in header for column in columns):
                        raise ReportError('table columns must be unique existing CSV headers')
                    indices = [header.index(column) for column in columns]
                    rows = []
                    for row in reader:
                        if len(row) != len(header):
                            raise ReportError('CSV row length does not match its header')
                        rows.append('<tr>' + ''.join('<td>' + escape(row[index]) + '</td>' for index in indices) + '</tr>')
                    caption = '<caption>' + escape(text(block['caption'], 'table caption')) + '</caption>' if 'caption' in block else ''
                    blocks.append('<div class="table-wrap"><table>' + caption + '<thead><tr>' + ''.join('<th scope="col">' + escape(column) + '</th>' for column in columns) + '</tr></thead><tbody>' + ''.join(rows) + '</tbody></table></div>')
                elif kind == 'figure':
                    fields(block, ('type', 'path', 'caption'))
                    caption = escape(text(block['caption'], 'figure caption'))
                    source = figure_data(block['path'], workspace.read(block['path'], location=location + '.path'))
                    figures += 1
                    blocks.append(f'<figure id="figure-{figures}"><img src="{source}" alt="{caption}"><figcaption>图 {figures} · {caption}</figcaption></figure>')
                else:
                    raise ReportError('unknown report block type')
            sections.append(f'<section id="{anchor}"><h2>{number}. {heading}</h2>' + ''.join(blocks) + '</section>')
        html = ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1">'
            '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'">'
            '<title>' + title + '</title><style>' + CSS + '</style></head><body><main><header><h1>' + title + '</h1>' + subtitle + '</header>'
            '<nav aria-label="目录"><strong>目录</strong><ol>' + ''.join(navigation) + '</ol></nav>' + ''.join(sections)
            + '<footer>本文由保存的文字、CSV 与图片排版生成。排版及源文件绑定不表示已验证数值口径、图注或结论。</footer></main></body></html>')
        workspace.write(output, html.encode('utf-8'))
        return {'schema': 'capstone-report-receipt/1', 'output': output, 'sections': len(sections),
            'figures': figures, 'source_sha256': workspace.sources}
    finally:
        os.close(workspace.root)


def render_presentation(spec_path, output):
    """Lay out bounded saved content. Never calculate, interpret or execute it."""
    try:
        from PIL import Image
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError:
        raise ReportError('PPTX output requires the verified python-pptx and Pillow environment') from None

    workspace = Workspace()
    try:
        spec = json.loads(workspace.read(spec_path, maximum=1024 * 1024))
        fields(spec, ('schema', 'title', 'slides'), ('subtitle',))
        if spec['schema'] != 'capstone-presentation/1':
            raise ReportError('PPTX requires capstone-presentation/1')
        if not isinstance(spec['slides'], list) or not 1 <= len(spec['slides']) <= 40:
            raise ReportError('presentation slides must contain 1..40 content records')
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        font = 'Noto Sans CJK SC'
        ink, teal, muted = RGBColor(23,43,59), RGBColor(33,119,120), RGBColor(82,103,116)
        footers, image_bytes, figures, tables, table_rows = [], 0, 0, 0, 0

        def literal(value, field, maximum):
            value = text(value, field)
            if len(value) > maximum or any(ord(c) < 32 and c != '\n' for c in value):
                raise ReportError(field + ' exceeds its text budget or contains control characters')
            return value

        def wrap(value, width):
            lines = []
            for line in value.split('\n'):
                current, units = '', 0
                for char in line:
                    weight = 0 if unicodedata.combining(char) else 2 if (ord(char)>127 or char.isupper() or char in 'mw@%#&') else 1
                    if current and units + weight > width:
                        lines.append(current);current, units = '', 0
                    current += char;units += weight
                lines.append(current)
            return lines

        def textbox(slide, x, y, width, height, lines, size, color=ink, bold=False):
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            frame = shape.text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = 0
            frame.margin_top = frame.margin_bottom = 0
            for index, line in enumerate(lines):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name, paragraph.font.size = font, Pt(size)
                paragraph.font.color.rgb, paragraph.font.bold = color, bold
                paragraph.line_spacing = 1.2
                paragraph.space_after = Pt(0)
            return shape

        def page(title, notes, *, cover=False):
            if len(deck.slides) >= 120:
                raise ReportError('presentation exceeds the 120 rendered-slide budget; prepare a smaller content spec')
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            lines = wrap(literal(title, 'slide title', 160), 38 if cover else 52)
            if len(lines) > (4 if cover else 2):
                raise ReportError('slide title is too long for its fixed layout')
            textbox(slide,.8,1.3 if cover else .35,11.7,2.3 if cover else .95,
                    lines,34 if cover else 27,teal,True)
            if notes:
                slide.notes_slide.notes_text_frame.text = '\n'.join(notes)
            footers.append((slide,notes))
            return slide

        cover = page(literal(spec['title'], 'title', 160), [], cover=True)
        if 'subtitle' in spec:
            lines = wrap(literal(spec['subtitle'], 'subtitle', 500),84)
            if len(lines)>6:raise ReportError('subtitle is too long for the cover layout')
            textbox(cover,.8,4,11.7,2,lines,18,muted)

        for index, content in enumerate(spec['slides']):
            fields(content, ('title','blocks'), ('source_notes',))
            title = literal(content['title'],'slide title',160)
            blocks = content['blocks']
            notes = content.get('source_notes',[])
            if not isinstance(notes,list) or len(notes)>4:
                raise ReportError('source_notes must contain at most four literal notes')
            notes = [literal(note,'source note',240) for note in notes]
            if len(wrap(' · '.join(notes),140))>2:
                raise ReportError('source_notes exceed the two-line footer budget')
            if not isinstance(blocks,list) or len(blocks)>20:
                raise ReportError('each content record allows at most 20 blocks')
            pending = []
            emitted = False
            def flush():
                nonlocal pending, emitted
                if not pending:return
                slide = page(title,notes)
                textbox(slide,.8,1.55,11.7,4.85,pending,18)
                pending=[];emitted=True
            def append_lines(lines):
                for line in lines:
                    if len(pending)>=14:flush()
                    pending.append(line)

            for block_index, block in enumerate(blocks):
                if not isinstance(block,dict):raise ReportError('presentation blocks must be objects')
                kind = block.get('type')
                location = f'slides[{index}].blocks[{block_index}]'
                if kind == 'paragraph':
                    fields(block,('type','text'))
                    append_lines(wrap(literal(block['text'],'paragraph',20000),74))
                elif kind == 'bullets':
                    fields(block,('type','items'))
                    if not isinstance(block['items'],list) or not 1<=len(block['items'])<=100:
                        raise ReportError('bullet items must contain 1..100 strings')
                    for item in block['items']:
                        lines=wrap(literal(item,'bullet',2000),70)
                        append_lines(['• '+lines[0]]+['  '+line for line in lines[1:]])
                elif kind == 'table':
                    flush()
                    fields(block,('type','csv','max_rows'),('columns','caption'))
                    limit=block['max_rows']
                    if type(limit) is not int or not 1<=limit<=1000:
                        raise ReportError('table max_rows must be an integer within 1..1000')
                    if not isinstance(block['csv'],str) or Path(block['csv']).suffix.lower()!='.csv':
                        raise ReportError('table sources must be saved CSV files')
                    data=workspace.read(block['csv'],location=location+'.csv',maximum=16*1024*1024)
                    csv.field_size_limit(65536)
                    reader=csv.reader(io.StringIO(data.decode('utf-8-sig'),newline=''),strict=True)
                    header=next(reader,[])
                    if not header or any(not name for name in header) or len(set(header))!=len(header):
                        raise ReportError('CSV must have a nonempty unique header')
                    columns=block.get('columns',header)
                    if (not isinstance(columns,list) or not 1<=len(columns)<=8
                            or any(not isinstance(name,str) or name not in header for name in columns)
                            or len(set(columns))!=len(columns)):
                        raise ReportError('select 1..8 unique existing CSV columns for the slide')
                    indices=[header.index(name) for name in columns]
                    rows=[]
                    for row in reader:
                        if len(rows)>=limit:raise ReportError('CSV exceeds explicit table max_rows; prepare a bounded display CSV')
                        if len(row)!=len(header):raise ReportError('CSV row length does not match its header')
                        rows.append([row[column] for column in indices])
                    caption=literal(block['caption'],'table caption',240) if 'caption' in block else ''
                    if len(wrap(caption,110))>2:raise ReportError('table caption exceeds its layout budget')
                    unit_width=max(8,int(11.7*72/len(columns)/7))
                    def height(row):
                        if any(len(value)>2000 or any(ord(c)<32 and c!='\n' for c in value) for value in row):
                            raise ReportError('CSV cell exceeds its literal text budget')
                        return max(24, max(len(wrap(value,unit_width)) for value in row)*15+8)
                    header_height=height(columns)
                    if header_height>140:raise ReportError('CSV headers are too long for the table layout')
                    pages, batch, used = [], [], header_height
                    for row in rows:
                        row_height=height(row)
                        if row_height+header_height>330:raise ReportError('CSV cell cannot fit on one slide; prepare a compact display CSV')
                        if used+row_height>330:
                            pages.append(batch);batch=[];used=header_height
                        batch.append((row,row_height));used+=row_height
                    pages.append(batch)
                    for batch in pages:
                        slide=page(title,notes);emitted=True
                        if caption:textbox(slide,.8,1.4,11.7,.5,wrap(caption,110),12,muted)
                        shape=slide.shapes.add_table(len(batch)+1,len(columns),Inches(.8),Inches(2),Inches(11.7),Pt(header_height+sum(h for _,h in batch)))
                        table=shape.table
                        for row_number,(values,row_height) in enumerate([(columns,header_height),*batch]):
                            table.rows[row_number].height=Pt(row_height)
                            for column,value in enumerate(values):
                                cell=table.cell(row_number,column)
                                cell.text=value
                                cell.margin_left=cell.margin_right=Inches(.06)
                                cell.margin_top=cell.margin_bottom=Inches(.03)
                                cell.text_frame.word_wrap=True
                                for paragraph in cell.text_frame.paragraphs:
                                    paragraph.font.name,paragraph.font.size=font,Pt(12)
                                    paragraph.font.bold=row_number==0
                                    paragraph.font.color.rgb=teal if row_number==0 else ink
                        tables+=1
                    table_rows+=len(rows)
                elif kind == 'figure':
                    flush();fields(block,('type','path','caption'))
                    path=block['path']
                    if not isinstance(path,str) or Path(path).suffix.lower() not in {'.png','.jpg','.jpeg'}:
                        raise ReportError('PPTX figures must be saved PNG or JPEG files')
                    data=workspace.read(path,location=location+'.path',maximum=8*1024*1024)
                    image_bytes+=len(data)
                    if image_bytes>32*1024*1024:raise ReportError('presentation figures exceed the 32 MiB processing budget')
                    try:
                        with Image.open(io.BytesIO(data)) as image:
                            width,height=image.size
                            if image.format not in {'PNG','JPEG'} or width*height>20000000 or not 1<=width<=20000 or not 1<=height<=20000:
                                raise ReportError('figure type or pixel budget is invalid')
                            image.verify()
                    except (OSError,ValueError,Image.DecompressionBombError):
                        raise ReportError('figure is not a valid bounded PNG or JPEG') from None
                    caption=wrap(literal(block['caption'],'figure caption',400),110)
                    if len(caption)>4:raise ReportError('figure caption exceeds its layout budget')
                    slide=page(title,notes);emitted=True
                    scale=min(11.7/width,4.3/height)
                    image_width,image_height=width*scale,height*scale
                    slide.shapes.add_picture(io.BytesIO(data),Inches(.8+(11.7-image_width)/2),Inches(1.45),
                        width=Inches(image_width),height=Inches(image_height))
                    textbox(slide,.8,5.9,11.7,.85,caption,12,muted)
                    figures+=1
                else:
                    raise ReportError('unknown presentation block type')
            flush()
            if not emitted:page(title,notes)
        for number,(slide,notes) in enumerate(footers,1):
            if notes:textbox(slide,.8,6.85,10.7,.4,wrap(' · '.join(notes),140),10,muted)
            textbox(slide,11.8,7.0,.7,.25,[f'{number}/{len(deck.slides)}'],10,muted)
        stream=io.BytesIO();deck.save(stream)
        if stream.tell()>64*1024*1024:raise ReportError('PPTX exceeds the 64 MiB output processing budget')
        workspace.write(output,stream.getvalue())
        return {'schema':'capstone-presentation-receipt/1','output':output,'slides':len(deck.slides),
            'figures':figures,'tables':tables,'table_rows':table_rows,'source_sha256':workspace.sources,
            'interpretation':'none; layout and source binding do not verify facts or conclusions'}
    finally:
        os.close(workspace.root)


def main():
    parser = argparse.ArgumentParser(description='Render report HTML or content-only PPTX JSON with literal text, exact CSV values and saved figures. HTML supports safe SVG/PNG/JPEG; PPTX supports PNG/JPEG. Workspace paths only; no raw HTML, scripts or network resources.')
    parser.add_argument('--version', action='version', version='capstone-report ' + VERSION)
    parser.add_argument('--spec', required=True, help='HTML: capstone-report/1 title, sections[{id,title,blocks}]. PPTX: capstone-presentation/1 title, subtitle?, slides[{title,blocks,source_notes?:[text]}]. Blocks: paragraph(text), bullets(items), table(csv,columns?,caption?,max_rows required for PPTX), figure(path,caption).')
    parser.add_argument('--output', required=True, help='HTML or .pptx output within the current workspace; parent directory must exist.')
    args = parser.parse_args()
    try:
        receipt = render(args.spec, args.output)
    except ReportError as error:
        print('capstone-report: ' + str(error), file=sys.stderr)
        if error.diagnostic:
            print(json.dumps(error.diagnostic, ensure_ascii=False), file=sys.stderr)
        return 2
    except (OSError, ValueError, csv.Error):
        print('capstone-report: cannot read or write a workspace file, or invalid UTF-8/JSON/CSV; check saved sources and paths', file=sys.stderr)
        return 2
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
