"""Real PPTX bytes, layout bounds, exact source strings and confined paths."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
import pytest

SCRIPT = Path(__file__).parents[2] / 'tools/native_report.py'


def spec(blocks, **extra):
    return {'schema':'capstone-presentation/1','title':'通用研究汇报',
        'slides':[{'title':'已保存的观察','blocks':blocks,'source_notes':['Source dataset revision r1']}],**extra}


def run(directory, value, output='report.pptx'):
    (directory/'slides.json').write_text(json.dumps(value,ensure_ascii=False))
    return subprocess.run([sys.executable,str(SCRIPT),'--spec','slides.json','--output',output],
        cwd=directory,capture_output=True,text=True)


def deck_text(deck):
    return '\n'.join(shape.text if shape.has_text_frame else '\n'.join(cell.text for row in shape.table.rows for cell in row.cells)
        if shape.has_table else '' for slide in deck.slides for shape in slide.shapes)


def test_real_pptx_has_exact_csv_strings_image_notes_and_source_hashes(tmp_path):
    (tmp_path/'data.csv').write_text('label,value,ratio\n甲省,001.2300,12.50%\n乙省,未记录,1.00%\n')
    Image.new('RGB',(320,180),'teal').save(tmp_path/'figure.png')
    value=spec([{'type':'paragraph','text':'仅排版已有观察；<script>是普通文字。'},
        {'type':'table','csv':'data.csv','max_rows':20,'caption':'保存的原始字符串'},
        {'type':'figure','path':'figure.png','caption':'已保存图片，不新增结论'}])
    result=run(tmp_path,value)
    assert result.returncode==0,result.stderr
    output=tmp_path/'report.pptx'
    assert output.read_bytes().startswith(b'PK')
    deck=Presentation(output)
    assert len(deck.slides)>=4
    text=deck_text(deck)
    assert all(item in text for item in ['001.2300','12.50%','未记录','Source dataset revision r1','<script>'])
    assert sum(shape.shape_type==13 for slide in deck.slides for shape in slide.shapes)==1
    assert all('Source dataset revision r1' in slide.notes_slide.notes_text_frame.text for slide in list(deck.slides)[1:])
    receipt=json.loads(result.stdout)
    assert receipt['schema']=='capstone-presentation-receipt/1' and receipt['slides']==len(deck.slides)
    for path,digest in receipt['source_sha256'].items():assert hashlib.sha256((tmp_path/path).read_bytes()).hexdigest()==digest
    with ZipFile(output) as archive:
        assert not any(b'TargetMode="External"' in archive.read(path) for path in archive.namelist() if path.endswith('.rels'))
    for slide in deck.slides:
        for shape in slide.shapes:
            assert 0<=shape.left and 0<=shape.top
            assert shape.left+shape.width<=deck.slide_width+100
            assert shape.top+shape.height<=deck.slide_height+100


def test_long_chinese_text_and_csv_rows_paginate_without_silent_loss(tmp_path):
    paragraphs='。'.join(f'第{i}项保留的说明与限制' for i in range(160))
    (tmp_path/'many.csv').write_text('name,value\n'+''.join(f'row-{i},00{i}.50%\n' for i in range(65)))
    result=run(tmp_path,spec([{'type':'paragraph','text':paragraphs},
        {'type':'table','csv':'many.csv','max_rows':65}]))
    assert result.returncode==0,result.stderr
    deck=Presentation(tmp_path/'report.pptx');text=deck_text(deck)
    assert len(deck.slides)>5 and '第159项' in text.replace('\n','') and 'row-64' in text and '0064.50%' in text
    tables=[shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    assert len(tables)>1 and all(table.cell(0,0).text=='name' for table in tables)
    assert sum(len(table.rows)-1 for table in tables)==65


@pytest.mark.parametrize('block', [
    {'type':'html','text':'<script>x</script>'},
    {'type':'table','csv':'../outside.csv','max_rows':5},
    {'type':'figure','path':'https://example.test/image.png','caption':'remote'},
    {'type':'figure','path':'active.svg','caption':'unsupported SVG'},
    {'type':'paragraph','text':'ok','script':'no'},
])
def test_arbitrary_code_remote_paths_and_unknown_shapes_fail_without_output(tmp_path,block):
    (tmp_path/'active.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg"><script>x</script></svg>')
    result=run(tmp_path,spec([block]))
    assert result.returncode==2 and not (tmp_path/'report.pptx').exists()


def test_csv_row_budget_is_explicit_and_never_silently_truncates(tmp_path):
    (tmp_path/'data.csv').write_text('name,value\na,1\nb,2\n')
    result=run(tmp_path,spec([{'type':'table','csv':'data.csv','max_rows':1}]))
    assert result.returncode==2 and 'max_rows' in result.stderr
    assert not (tmp_path/'report.pptx').exists()


def test_spec_size_and_source_symlinks_cannot_bypass_workspace_boundary(tmp_path):
    workspace=tmp_path/'workspace';workspace.mkdir()
    Image.new('RGB',(2,2),'red').save(tmp_path/'outside.png')
    (workspace/'link.png').symlink_to(tmp_path/'outside.png')
    assert run(workspace,spec([{'type':'figure','path':'link.png','caption':'symlink'}])).returncode==2
    assert run(workspace,spec([]),output='../outside.pptx').returncode==2
    assert run(workspace,spec([]),output='slides.json').returncode==2
    assert run(workspace,spec([{'type':'paragraph','text':'x'*1100000}])).returncode==2
    assert not (workspace/'report.pptx').exists()


@pytest.mark.parametrize('block', [
    {'type':'table','csv':'data.csv','max_rows':True},
    {'type':'table','csv':'data.csv'},
    {'type':'table','csv':'data.csv','max_rows':1001},
    {'type':'table','csv':'data.csv','max_rows':2,'columns':['unknown']},
    {'type':'table','csv':'data.json','max_rows':2},
])
def test_table_contract_rejects_ambiguous_or_unbounded_requests(tmp_path,block):
    (tmp_path/'data.csv').write_text('name,value\na,1\n')
    (tmp_path/'data.json').write_text('name,value\na,1\n')
    assert run(tmp_path,spec([block])).returncode==2
    assert not (tmp_path/'report.pptx').exists()


def test_source_processing_and_pixel_budgets_fail_before_output(tmp_path):
    with (tmp_path/'huge.csv').open('wb') as output:output.truncate(16*1024*1024+1)
    assert run(tmp_path,spec([{'type':'table','csv':'huge.csv','max_rows':10}])).returncode==2
    Image.new('1',(5000,5000)).save(tmp_path/'too-many-pixels.png')
    assert run(tmp_path,spec([{'type':'figure','path':'too-many-pixels.png','caption':'large'}])).returncode==2
    assert not (tmp_path/'report.pptx').exists()


def test_hardlink_output_cannot_replace_spec_or_saved_source(tmp_path):
    (tmp_path/'source.csv').write_text('name,value\na,1\n')
    (tmp_path/'report.pptx').hardlink_to(tmp_path/'source.csv')
    before=(tmp_path/'source.csv').read_bytes()
    result=run(tmp_path,spec([{'type':'table','csv':'source.csv','max_rows':10}]))
    assert result.returncode==2
    assert (tmp_path/'source.csv').read_bytes()==before


@pytest.mark.parametrize('extension', ['html','pptx'])
@pytest.mark.parametrize('source_kind', ['spec','csv','image'])
def test_input_hardlinks_cannot_import_external_spec_csv_or_image(tmp_path,extension,source_kind):
    workspace=tmp_path/'workspace';workspace.mkdir()
    if source_kind=='csv':
        source=tmp_path/'outside.csv';source.write_text('label,value\nexternal,001.25\n')
        (workspace/'data.csv').hardlink_to(source)
        blocks=[{'type':'table','csv':'data.csv',**({'max_rows':5} if extension=='pptx' else {})}]
    elif source_kind=='image':
        source=tmp_path/'outside.png';Image.new('RGB',(2,2),'red').save(source)
        (workspace/'figure.png').hardlink_to(source)
        blocks=[{'type':'figure','path':'figure.png','caption':'External hardlinked image'}]
    else:
        blocks=[{'type':'paragraph','text':'External hardlinked spec'}]
    value=spec(blocks) if extension=='pptx' else {'title':'Report','sections':[{'id':'results','title':'Results','blocks':blocks}]}
    if source_kind=='spec':
        source=tmp_path/'outside.json';source.write_text(json.dumps(value))
        (workspace/'slides.json').hardlink_to(source)
    else:
        (workspace/'slides.json').write_text(json.dumps(value))
    before=source.read_bytes()
    result=subprocess.run([sys.executable,str(SCRIPT),'--spec','slides.json','--output','report.'+extension],
        cwd=workspace,capture_output=True,text=True)
    assert result.returncode==2 and 'hardlinked' in result.stderr
    assert not (workspace/('report.'+extension)).exists()
    assert source.read_bytes()==before


def test_build_smoke_checks_real_pptx_cli_and_sources(tmp_path):
    import importlib.util
    loaded=importlib.util.spec_from_file_location('native_analysis_smoke',SCRIPT.with_name('native_analysis_smoke.py'))
    module=importlib.util.module_from_spec(loaded);loaded.loader.exec_module(module)
    module.exercise_presentation(tmp_path,renderer=[sys.executable,str(SCRIPT)])
    deck=Presentation(tmp_path/'presentation.pptx')
    assert len(deck.slides)>=4 and '001.2300' in deck_text(deck)
